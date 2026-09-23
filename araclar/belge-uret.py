#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
USTAD-MOTOR / belge-uret.py
Renkli, biçimli Word ana plan belgesi üretir: Masaüstü/USTAD-MOTOR-PROJE-ANA-PLANI.docx
Kenan biçemli Word ister: renkli başlık şeritleri, tablolar, bölüm renkleri.
"""
import os
from datetime import date
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HEDEF = r"C:\Users\kenan\OneDrive\Desktop\USTAD-MOTOR-PROJE-ANA-PLANI.docx"

RENKLER = {
    "ana": "0E9F8E", "program": "4F46E5", "oneriyor": "E08A00", "testler": "12A150",
    "deneme": "DC2626", "notlar": "7C3AED", "istatistik": "1D6FE0", "ayarlar": "556070",
    "koyu": "1B2A6B", "gri": "EEF2F7"
}


def golge(parca, renk):
    """Paragrafa arka plan rengi verir."""
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), renk)
    parca._p.get_or_add_pPr().append(shd)


def serit(doc, metin, renk_kod):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("  " + metin)
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    golge(p, renk_kod)
    return p


def yazi(doc, metin, kalin=False, boyut=10.5, italik=False):
    p = doc.add_paragraph()
    r = p.add_run(metin)
    r.bold = kalin
    r.italic = italik
    r.font.size = Pt(boyut)
    p.paragraph_format.space_after = Pt(4)
    return p


def madde(doc, metin, renk_kod=None):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(metin)
    r.font.size = Pt(10.5)
    p.paragraph_format.space_after = Pt(2)
    return p


def tablo(doc, basliklar, satirlar, renk_kod, genislikler=None):
    t = doc.add_table(rows=1, cols=len(basliklar))
    t.style = "Table Grid"
    hdr = t.rows[0].cells
    for i, b in enumerate(basliklar):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        r = p.add_run(b)
        r.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        golge(p, renk_kod)
        satir_golge(hdr[i], renk_kod)
    for j, satir in enumerate(satirlar):
        hucreler = t.add_row().cells
        for i, deger in enumerate(satir):
            hucreler[i].text = ""
            p = hucreler[i].paragraphs[0]
            r = p.add_run(str(deger))
            r.font.size = Pt(10)
            if j % 2 == 1:
                golge(p, "F7FAFC")
                satir_golge(hucreler[i], "F7FAFC")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def satir_golge(hucre, renk):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), renk)
    hucre._tc.get_or_add_tcPr().append(shd)


def kapak(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("ÜSTAD MOTOR")
    r.bold = True
    r.font.size = Pt(34)
    r.font.color.rgb = RGBColor(0x1B, 0x2A, 0x6B)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("SINAV KOÇLUĞU PLATFORMU  ·  PROJE ANA PLANI")
    r.font.size = Pt(12)
    r.bold = True
    r.font.color.rgb = RGBColor(0x0E, 0x9F, 0x8E)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Tek motor, sınav başına ayrı uygulama (APK)   •   Hazırlayan: ÜSTAD   •   " +
                  date.today().strftime("%d.%m.%Y"))
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0x55, 0x60, 0x70)
    logo = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                        "tasarim", "ustad-logo-tam.png")
    if os.path.exists(logo):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(logo, width=Cm(9))


def main():
    doc = Document()
    stil = doc.styles["Normal"]
    stil.font.name = "Segoe UI"
    stil.font.size = Pt(10.5)
    bolum = doc.sections[0]
    bolum.left_margin = bolum.right_margin = Cm(2)

    kapak(doc)

    serit(doc, "1. AMAÇ", RENKLER["ana"])
    yazi(doc, "ÜSTAD, öğrencinin sınava hazırlığını baştan sona yöneten bir çalışma koçudur. "
              "Öğrenciyi tanır, adıyla hitap eder, kişiye özel ders programı kurar, uyumunu takip eder, "
              "yanlışlarından eksik konusunu bulur ve o konuya geri götürür. Amaç: hazırlığı ölçülebilir, "
              "takip edilebilir ve kişiye özel kılmak.")
    yazi(doc, "Bu belge, Kenan'ın ÜSTAD.docx fikir dokümanının teknik karşılığıdır; "
              "karara bağlanan maddeler, uygulama sırası ve yapım kuralları buradadır.", italik=True)

    serit(doc, "2. MİMARİ: TEK MOTOR + SINAV PAKETİ", RENKLER["program"])
    yazi(doc, "Bütün sınavlar aynı motoru kullanır. Sınava özel her şey paketin içindedir. "
              "Böylece motorda yapılan bir düzeltme bütün uygulamalara aynı anda yarar, "
              "yeni sınav eklemek yeni paket yazmakla kalır.")
    tablo(doc, ["Katman", "İçerik", "Her sınavda"],
          [["MOTOR (kod)", "Ekranlar, menü, test/deneme mantığı, program üretici, istatistik, bildirim, sesli karşılama",
            "Aynı kalır"],
           ["PAKET (içerik)", "Sınav adı, ders listesi, konu listesi, soru bankası, ders notları, deneme takvimi, duyuru beslemesi",
            "Sınava göre değişir"],
           ["KABUK (Android)", "WebView, ikon, uygulama adı, imza, bildirim izinleri",
            "Sınav adı ve ikonu değişir"]],
          RENKLER["program"])
    yazi(doc, "Sonuç: ÜSTAD KPSS-B, ÜSTAD Ön Lisans, ÜSTAD AGS, ÜSTAD TYT … ayrı ayrı kurulabilen "
              "uygulamalar olur; hepsi aynı kalitede ilerler.", kalin=True)

    serit(doc, "3. MENÜ VE MODÜLLER", RENKLER["oneriyor"])
    tablo(doc, ["Menü", "Ne yapar", "Kural"],
          [["Ana Sayfa", "Sistemi anlatır, uygulamanın nasıl verimli kullanılacağını söyler",
            "Sağda akan DUYURULAR panosu; 'Şimdi Oluştur' düğmesi program ekranına götürür"],
           ["Ders Programı", "7 soruluk anket → kişiye özel haftalık program", "Zorlanan derslerin arasına sevilen dersler konur; konu çalışması, soru çözümü ve boş zaman blokları yazılır"],
           ["ÜSTAD Öneriyor", "Yanlışlardan eksik konu tespiti", "'Eksik Konuya Git' düğmesi; tekrar sonrası en az 5 telafi sorusu"],
           ["Testler", "Konu testleri", "30 soru, SÜRESİZ; 'Boş Bıraktığın' ve 'Yanlış Cevaplanan' düğmeleri; 'Soru ile İlgili Konuya Git'"],
           ["Deneme Sınavı", "Gerçek sınav formatında deneme", "Süre giriş anında başlar; süre bitince 15 dakika erişim kilidi; öğrenci geçtiği soruları görebilir; her hafta 1 deneme"],
           ["Ders Notları", "Sınav konusunu kapsayan ders notları", "Her dersin yanında 'Ders Notunu Oku' düğmesi"],
           ["İstatistik", "Çözülen soru, doğru/yanlış, başarı, uyum", "Program uyumu 'uyulmadı' olarak işlenir"],
           ["Ayarlar", "Renk paleti, düzen, sesli karşılama", "Her bölümün ayrı rengi korunur"]],
          RENKLER["oneriyor"])
    yazi(doc, "Menü, sol üstteki üç çizgide (☰) durur ve bütün bölümlerde görünür.", kalin=True)

    serit(doc, "4. KPSS-B KAPSAMI (yalnız o sınavın dersleri)", RENKLER["testler"])
    tablo(doc, ["Bölüm", "Ders", "Soru"],
          [["Genel Yetenek", "Türkçe", "30"], ["Genel Yetenek", "Matematik", "30"],
           ["Genel Kültür", "Tarih", "27"], ["Genel Kültür", "Coğrafya", "18"],
           ["Genel Kültür", "Vatandaşlık", "9"], ["Genel Kültür", "Güncel Bilgiler", "6"],
           ["TOPLAM", "Genel Yetenek + Genel Kültür", "120 soru / 130 dakika"]],
          RENKLER["testler"])
    yazi(doc, "Kapsam kuralı: KPSS paketine TYT/AYT/YDS dersleri girmez. Her sınavın dersleri kendi paketinde yaşar.", italik=True)

    serit(doc, "5. SINAV SIRASI (tek tek ele alınacak)", RENKLER["deneme"])
    tablo(doc, ["Sıra", "Uygulama", "Durum"],
          [["1", "ÜSTAD KPSS-B (Lisans)", "Şu an yapılıyor"],
           ["2", "ÜSTAD KPSS Ön Lisans / Ortaöğretim", "Aynı paket; ad ve zorluk farkı"],
           ["3", "ÜSTAD KPSS-A", "Alan/uzmanlık yapısı ayrı paket"],
           ["4", "ÜSTAD MEB-AGS / ÖABT", "Kılavuz doğrulaması sonrası"],
           ["5", "ÜSTAD TYT / AYT / YDT", "Çok bölümlü sınav yapısı; en son"],
           ["—", "ÜSTAD LGS", "Not: dokümandaki 'LYS' ifadesi LGS olarak düzeltildi (LYS 2018'de kaldırıldı, lise girişi LGS'dir)"]],
          RENKLER["deneme"])

    serit(doc, "6. KARARA BAĞLANAN KONULAR", RENKLER["koyu"])
    tablo(doc, ["Konu", "Karar", "Kim verdi"],
          [["Lisans sunucusu (15 gün deneme, tek cihaz)", "ŞİMDİLİK YAPILMAYACAK — uygulama tanıtım amaçlı çıkacak, lisans sonraki iş",
            "Kenan"],
           ["Logo", "Dokümandaki PNG kullanılacak; bizim ürettiğimiz şeffaf sürümler beğenilmezse PNG'ye dönülür", "Kenan"],
           ["Eski soru bankası", "İsraf yok: mevcut USTAD-KPSS soruları yeni motora taşınır", "Kenan"],
           ["Sesli karşılama", "Cihazın Türkçe metin-okuma motoru kullanılır (ücretsiz, çevrimdışı)", "Teknik"],
           ["Duyurular panosu", "Haber servisi (RSS→JSON) bağlanana kadar örnek besleme; başlığa dokununca haber uygulama İÇİNDE açılır", "Teknik"],
           ["Bildirim", "Ders saatinden 5 dakika önce; Android 12+ için tam zamanlı alarm ve bildirim izni istenir", "Teknik"]],
          RENKLER["koyu"])

    serit(doc, "7. TELİF VE YASAL ÇERÇEVE", RENKLER["notlar"])
    yazi(doc, "ÖSYM'nin geçmiş sınav soruları telif korumalıdır; birebir kullanılamaz. Bu yüzden:")
    madde(doc, "Geçmiş sınavların konu dağılımı, soru sayısı ve soru formatı analiz edilir (dağılım taklit edilir).")
    madde(doc, "Sorular ÖZGÜN olarak yazılır; hiçbir soru kopyalanmaz.")
    madde(doc, "Her pakete yasal uyarı sayfası konur (daha önce yapıldığı gibi).")
    madde(doc, "Haber başlıkları kaynak belirtilerek ve kısa alıntı olarak gösterilir.")

    serit(doc, "8. İÇERİK ÜRETİM FAZLARI", RENKLER["istatistik"])
    tablo(doc, ["Faz", "İçerik", "Hedef"],
          [["Faz 1", "Tüm ders notları + her konunun ilk testi + 5 deneme", "Uygulama yayına hazır hâle gelir"],
           ["Faz 2", "Her konuya 3 test (30'ar soru) + deneme sayısı 20", "Konu hâkimiyeti ölçümü tamamlanır"],
           ["Faz 3", "Her hafta 1 deneme eklenir (sınav tarihine kadar)", "%100 sınav kapsamı"],
           ["Faz 4", "Güncel bilgiler haftalık bülten + duyuru beslemesi", "Canlı içerik"]],
          RENKLER["istatistik"])
    yazi(doc, "Gerçek hacim: 100 hafta × 120 soru = 12.000 deneme sorusu, artı konu testleri. "
              "Bu yüzden içerik paketi sunucudan güncellenebilir yapılır; uygulamayı yeniden kurmak gerekmez.", kalin=True)

    serit(doc, "9. BUGÜNE KADAR ÜRETİLENLER", RENKLER["ayarlar"])
    tablo(doc, ["Dosya", "Ne işe yarar"],
          [["USTAD-MOTOR\\index.html", "Uygulamanın iskeleti: karşılama, 7 bölüm, test/deneme modülü"],
           ["assets\\stil.css", "12 renk paleti, bölüm renkleri, bütün biçim"],
           ["assets\\motor.js", "Motorun çekirdeği: menü, program üretici, test mantığı, istatistik, sesli karşılama, haber penceresi"],
           ["assets\\deneme.js", "Deneme sınavı motoru: süre sayacı, 15 dakika kilidi, optik form, ders bazlı sonuç"],
           ["icerik\\veri.js", "KPSS-B paketi: 6 ders, 273 konu başlığı, deneme takvimi, ders notu dizini"],
           ["icerik\\sorular.js", "Derlenmiş soru bankası: 120 özgün soru (Türkçe 30, Matematik 30, Tarih 27, Coğrafya 18, Vatandaşlık 9, Güncel 6)"],
           ["icerik\\notlar.js", "20 tam ders notu (konu anlatımı + püf noktaları + sınav ipucu)"],
           ["icerik\\besleme.js", "Canlı haber beslemesi: 6 kaynak, 40 güncel eğitim/sınav haberi"],
           ["araclar\\logo-isle.py", "Dokümandaki logoyu şeffaflaştırır, koyu zemin sürümünü ve kafa işaretini üretir"],
           ["araclar\\haber-cek.py", "Haber kaynaklarını çeker, süzer, tekrarları atar, besleme dosyasını yazar"],
           ["araclar\\sorulari-derle.py", "Soru/not dosyalarını derler ve kalite denetimi yapar (şık dağılımı, tekrar, konu kapsaması)"],
           ["araclar\\ekran-cek.sh", "20 ekranın görüntüsünü otomatik alır (kanıt klasörü)"],
           ["ekran-goruntuleri\\", "Alınan ekran görüntüleri (01-giriş … 20-ders-notları)"]],
          RENKLER["ayarlar"])

    serit(doc, "10. İÇERİK DURUMU (bu tarihte)", RENKLER["testler"])
    tablo(doc, ["Kalem", "Hedef (tam KPSS-B)", "Şu an", "Durum"],
          [["Deneme sorusu (tam bir deneme seti)", "120", "120", "İlk tam deneme hazır"],
           ["Ders notu", "273 konu başlığı", "20 tam konu", "Devam ediyor"],
           ["Konu testi sorusu", "273 konu × 30 soru", "120 (deneme ile ortak)", "Devam ediyor"],
           ["Canlı haber kaynağı", "6 kaynak", "6 kaynak / 40 haber", "Çalışıyor"],
           ["Soru kalite denetimi", "şık dağılımı + tekrar + kapsam", "Otomatik denetim kurulu", "Çalışıyor"]],
          RENKLER["testler"])
    yazi(doc, "Not: yazılan sorular özgün ve otomatik denetimden geçmiştir. Yayından önce her sorunun " +
              "gözle son kontrolü yapılmalıdır (matematik çözümleri ve tarih/anayasa bilgileri örneklemeli olarak elle doğrulandı).",
         italik=True)

    serit(doc, "11. SIRADAKİ ADIMLAR", RENKLER["ana"])
    madde(doc, "Ders notlarının 273 konu başlığına tamamlanması (Faz 1).")
    madde(doc, "Her konuya 30 soruluk test üretimi ve denetimi (Faz 2).")
    madde(doc, "Bildirim motorunun kurulması ve program uyum takibi.")
    madde(doc, "Android kabuğunun yazılması ve ilk APK'nın imzalanması (paket adı, ikon, sürüm).")
    madde(doc, "Yayın paketi: siteye yükleme, TV ve telefon sürümleri, güncelleme kontrolü.")

    doc.save(HEDEF)
    print("Yazıldı:", HEDEF, os.path.getsize(HEDEF) // 1024, "KB")


if __name__ == "__main__":
    main()
